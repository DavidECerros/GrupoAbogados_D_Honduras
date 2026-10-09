import io
import sqlite3

import pytest
from conftest import authenticate, create
from openpyxl import load_workbook

from backend import db as database


def preview_csv(c, entity, kind, rows, mode="create"):
    return c.post(
        f"/api/entities/{entity}/catalogs/{kind}/import/preview",
        files={"file": ("data.csv", rows.encode("utf-8"), "text/csv")},
        data={"mode": mode, "reason": "Importación de prueba"},
    )


def test_catalog_import_preview_commit_and_export(portfolio):
    c, e, customer, lot, sale = portfolio
    csv = "ID,Nombre,Tipo de documento,Documento,Teléfono,Correo,Observaciones,Estado\n,Persona nueva,DNI,000000123,0000123,,,Activo\n"
    result = preview_csv(c, e["id"], "clients", csv)
    assert result.status_code == 200, result.text
    plan = result.json()
    assert plan["create"] == 1 and plan["error_count"] == 0
    assert c.get(f"/api/entities/{e['id']}/clients").json()["total"] == 1
    saved = create(
        c, "/api/catalogs/import/execute", {"token": plan["token"], "confirmation": "IMPORTAR"}
    )
    assert saved["imported"] == 1 and saved["backup"].endswith(".zip")
    assert (
        c.post(
            "/api/catalogs/import/execute",
            json={"token": plan["token"], "confirmation": "IMPORTAR"},
        ).status_code
        == 409
    )
    exported = c.get(f"/api/entities/{e['id']}/catalogs/clients/export.xlsx")
    book = load_workbook(io.BytesIO(exported.content))
    assert book["Clientes"]["D3"].value == "000000123"
    assert book["Clientes"]["E3"].value == "0000123"
    body = c.post(
        f"/api/entities/{e['id']}/catalogs/clients/import/preview",
        files={"file": ("update.xlsx", exported.content)},
        data={"mode": "upsert", "reason": "Reimportar exportación"},
    )
    assert body.status_code == 200, body.text
    assert body.json()["update"] == 2 and body.json()["error_count"] == 0


def test_invalid_import_never_saves_and_duplicate_rows(portfolio):
    c, e, customer, lot, sale = portfolio
    csv = "ID,Código,Descripción,Precio,Moneda,Estado\n,B-01,Prueba,200,HNL,\n,B-01,Duplicado,200,HNL,\n,C-01,Inválido,-100,HNL,\n"
    result = preview_csv(c, e["id"], "lots", csv)
    assert result.status_code == 200, result.text
    assert result.json()["error_count"] == 2 and result.json()["token"] is None
    assert c.get(f"/api/entities/{e['id']}/lots").json()["total"] == 1


def test_import_updates_lot_base_price_not_contract(portfolio):
    c, e, customer, lot, sale = portfolio
    csv = f"ID,Código,Descripción,Precio,Moneda,Estado\n{lot['id']},A-01,Actualizado,220000,HNL,Vendido\n"
    plan = preview_csv(c, e["id"], "lots", csv, "upsert").json()
    assert plan["error_count"] == 0, plan
    create(c, "/api/catalogs/import/execute", {"token": plan["token"], "confirmation": "IMPORTAR"})
    assert c.get(f"/api/contracts/{sale['id']}").json()["price"] == 10_000_000
    assert c.get(f"/api/entities/{e['id']}/lots").json()["items"][0]["price"] == 22_000_000


def test_import_stale_and_cross_entity_id(portfolio):
    c, e, customer, lot, sale = portfolio
    csv = "ID,Código,Descripción,Precio,Moneda,Estado\n,B-01,Prueba,200,HNL,\n"
    plan = preview_csv(c, e["id"], "lots", csv).json()
    create(c, f"/api/entities/{e['id']}/lots", {"code": "C-01", "price": "100", "currency": "HNL"})
    assert (
        c.post(
            "/api/catalogs/import/execute",
            json={"token": plan["token"], "confirmation": "IMPORTAR"},
        ).status_code
        == 409
    )
    other = create(c, "/api/entities", {"name": "Otra entidad", "currency": "HNL"})
    bad = f"ID,Código,Descripción,Precio,Moneda,Estado\n{lot['id']},B-02,Prueba,200,HNL,\n"
    assert preview_csv(c, other["id"], "lots", bad, "upsert").json()["error_count"] == 1


def test_templates_are_downloadable_and_text_formatted(admin_client):
    c = admin_client
    for kind, sheet in [("clients", "Clientes"), ("lots", "Lotes")]:
        response = c.get(f"/api/catalogs/{kind}/template.xlsx")
        assert response.status_code == 200
        book = load_workbook(io.BytesIO(response.content))
        assert "Instrucciones" in book.sheetnames
        assert book[sheet]["B2"].number_format == "@"
        assert book[sheet]["B2"].value is None


def test_sql_dry_run_backup_commit_and_replay(portfolio):
    c, e, customer, lot, sale = portfolio
    sql = f"UPDATE clients SET phone='0000999' WHERE id={customer['id']};"
    plan = create(
        c, "/api/admin/database/sql/preview", {"sql": sql, "reason": "Corregir teléfono de prueba"}
    )
    assert plan["changes"] == 1 and plan["writes"]
    assert (
        c.get(f"/api/entities/{e['id']}/clients").json()["items"][0]["phone"] == customer["phone"]
    )
    result = create(
        c, "/api/admin/database/sql/execute", {"token": plan["token"], "confirmation": "EJECUTAR"}
    )
    assert result["changes"] == 1 and result["backup"].endswith(".zip")
    assert c.get(f"/api/entities/{e['id']}/clients").json()["items"][0]["phone"] == "0000999"
    assert (
        c.post(
            "/api/admin/database/sql/execute",
            json={"token": plan["token"], "confirmation": "EJECUTAR"},
        ).status_code
        == 409
    )
    with database.transaction() as db:
        assert db.execute("SELECT COUNT(*) FROM audit WHERE action='admin_sql'").fetchone()[0] == 1


@pytest.mark.parametrize(
    "sql",
    [
        "DROP TABLE clients",
        "PRAGMA foreign_keys=OFF",
        "DELETE FROM audit",
        "ATTACH DATABASE 'evil.db' AS evil",
        "SELECT load_extension('evil')",
        "UPDATE sessions SET csrf='x'",
        "SELECT 1; DELETE FROM clients",
    ],
)
def test_sql_cannot_bypass_schema_or_audit(portfolio, sql):
    c, *_ = portfolio
    assert (
        c.post(
            "/api/admin/database/sql/preview", json={"sql": sql, "reason": "Prueba de límites"}
        ).status_code
        == 422
    )


def test_sql_stale_preview_and_read_query(portfolio):
    c, e, customer, lot, sale = portfolio
    read = create(
        c,
        "/api/admin/database/sql/preview",
        {"sql": "SELECT name,phone FROM clients", "reason": "Consultar clientes"},
    )
    assert not read["writes"] and read["token"] is None
    plan = create(
        c,
        "/api/admin/database/sql/preview",
        {"sql": "UPDATE clients SET phone='1000'", "reason": "Cambio temporal"},
    )
    c.patch(
        f"/api/clients/{customer['id']}/status",
        json={"status": "inactive", "reason": "Cambio concurrente"},
    )
    assert (
        c.post(
            "/api/admin/database/sql/execute",
            json={"token": plan["token"], "confirmation": "EJECUTAR"},
        ).status_code
        == 409
    )


def test_database_export_is_consistent_sqlite(portfolio, tmp_path):
    c, e, customer, lot, sale = portfolio
    response = c.get("/api/admin/database/export.sqlite3")
    assert response.status_code == 200 and response.content.startswith(b"SQLite format 3")
    path = tmp_path / "export.sqlite3"
    path.write_bytes(response.content)
    with sqlite3.connect(path) as db:
        assert db.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
        assert db.execute("SELECT name FROM clients").fetchone()[0] == customer["name"]


def test_operator_cannot_import_or_access_database(portfolio):
    c, e, customer, lot, sale = portfolio
    create(
        c,
        "/api/users",
        {
            "username": "opdata",
            "full_name": "Operador",
            "password": "ClaveSeguraPrueba123!",
            "entity_ids": [e["id"]],
        },
    )
    authenticate(c, "opdata")
    c.post(
        "/api/auth/password",
        json={
            "current_password": "ClaveSeguraPrueba123!",
            "new_password": "OperadorDefinitivo123!",
        },
    )
    assert c.get("/api/admin/database").status_code == 403
    assert c.get("/api/admin/database/export.sqlite3").status_code == 403
    assert (
        c.post(
            "/api/admin/database/sql/preview", json={"sql": "SELECT 1", "reason": "Prueba permisos"}
        ).status_code
        == 403
    )
    assert (
        preview_csv(
            c,
            e["id"],
            "lots",
            "ID,Código,Descripción,Precio,Moneda,Estado\n,B-01,Prueba,200,HNL,\n",
        ).status_code
        == 403
    )


def test_relocate_keeps_original_freezes_and_changes_config(portfolio, tmp_path, monkeypatch):
    c, e, customer, lot, sale = portfolio
    target = tmp_path / "new-data"
    config = tmp_path / "location.json"
    monkeypatch.delenv("GESTOR_DATA_DIR", raising=False)
    monkeypatch.setattr(database, "LOCATION_CONFIG", config)
    monkeypatch.setattr(database, "MAINTENANCE", False)
    original = database.DATA
    # Fixture DATA is itself tmp_path; use a sibling target so it is outside it.
    target = tmp_path.parent / (tmp_path.name + "-new-data")
    result = c.post(
        "/api/admin/database/location",
        json={"path": str(target), "reason": "Traslado de prueba", "confirmation": "MOVER"},
    )
    assert result.status_code == 200, result.text
    assert result.json()["restart_required"]
    assert (original / "gestor.sqlite3").is_file()
    assert (target / "gestor.sqlite3").is_file()
    assert database.configured_data() == target.resolve()
    assert c.get("/api/admin/database").status_code == 503
    with sqlite3.connect(target / "gestor.sqlite3") as db:
        assert db.execute("SELECT name FROM clients").fetchone()[0] == customer["name"]
