import io
import json
import sqlite3
from concurrent.futures import ThreadPoolExecutor
from datetime import date
from decimal import Decimal
from uuid import uuid4

from conftest import authenticate, create
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend import artifacts
from backend import db as database
from backend.app import app
from backend.finance import create_payment, create_sale
from backend.models import Payment, Sale


def payment(contract, amount="10000", currency="HNL", **kwargs):
    return {
        "contract_id": contract["id"],
        "amount": amount,
        "currency": currency,
        "operation_id": str(uuid4()),
        **kwargs,
    }


def test_unique_setup_and_no_anonymous_access(client):
    assert client.get("/api/entities").status_code == 401
    assert (
        client.post(
            "/api/setup", json={"username": "admin", "full_name": "Admin", "password": "short"}
        ).status_code
        == 422
    )
    data = {"username": "admin", "full_name": "Admin", "password": "PruebaSuperSegura123"}
    assert client.post("/api/setup", json=data).status_code == 200
    assert client.post("/api/setup", json=data).status_code == 409


def test_down_payment_balance_and_month_end(portfolio):
    c, e, cl, lot, sale = portfolio
    account = c.get(f"/api/contracts/{sale['id']}").json()
    assert account["balance"] == 10_000_000
    assert account["obligations"][2]["due_date"] == "2026-02-28"
    assert account["obligations"][3]["due_date"] == "2026-03-31"
    paid = create(c, "/api/payments", payment(sale))
    assert paid["balance_after"] == 9_000_000
    assert c.get(f"/api/contracts/{sale['id']}").json()["balance"] == 9_000_000


def test_idempotency_allocation_and_excess(portfolio):
    c, e, cl, lot, sale = portfolio
    body = payment(sale, "15000")
    a = create(c, "/api/payments", body)
    b = create(c, "/api/payments", body)
    assert a["id"] == b["id"] and a["folio"] == b["folio"]
    assert c.post("/api/payments", json={**body, "amount": "1"}).status_code == 409
    assert c.post("/api/payments", json=payment(sale, "85000.01")).status_code == 409
    with database.transaction() as db:
        amounts = [
            r[0]
            for r in db.execute("SELECT amount FROM allocations WHERE payment_id=?", (a["id"],))
        ]
        assert sum(amounts) == 1_500_000
        assert amounts[0] == 1_000_000
        assert db.execute("SELECT COUNT(*) FROM payments").fetchone()[0] == 1


def test_currency_conversion_and_historical_rate(admin_client):
    c = admin_client
    e = create(c, "/api/entities", {"name": "USD", "currency": "USD"})
    cl = create(
        c,
        f"/api/entities/{e['id']}/clients",
        {"name": "Persona USD", "document_type": "DNI", "document": "TESTUSD", "phone": "000"},
    )
    lot = create(
        c, f"/api/entities/{e['id']}/lots", {"code": "U1", "price": "1000", "currency": "USD"}
    )
    sale = create(
        c,
        f"/api/entities/{e['id']}/contracts",
        {
            "lot_id": lot["id"],
            "client_id": cl["id"],
            "price": "1000",
            "currency": "USD",
            "down_payment": "0",
            "installments": 3,
            "down_due": "2026-01-01",
            "first_due": "2026-01-31",
        },
    )
    assert c.put("/api/settings", json={"exchange_rate": "25"}).status_code == 200
    paid = create(c, "/api/payments", payment(sale, "2500", "HNL"))
    assert paid["applied"] == 10000 and paid["rate"] == "25"
    assert c.put("/api/settings", json={"exchange_rate": "30"}).status_code == 200
    with database.transaction() as db:
        assert json.loads(db.execute("SELECT snapshot FROM receipts").fetchone()[0])["rate"] == "25"
    assert c.get(f"/api/contracts/{sale['id']}").json()["balance"] == 90000


def test_pdf_failure_retry_and_void(portfolio, monkeypatch):
    c, e, cl, lot, sale = portfolio
    original = artifacts.generate_receipt

    def failure(*args):
        raise OSError("fallo simulado")

    monkeypatch.setattr(artifacts, "generate_receipt", failure)
    paid = create(c, "/api/payments", payment(sale))
    assert paid["receipt_status"] == "pending" and "receipt_warning" in paid
    monkeypatch.setattr(artifacts, "generate_receipt", original)
    assert c.post(f"/api/receipts/{paid['receipt_id']}/regenerate").status_code == 200
    pdf = c.get(f"/api/receipts/{paid['receipt_id']}/pdf")
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    assert (
        c.post(
            f"/api/payments/{paid['id']}/void", json={"reason": "Cobro de prueba anulado"}
        ).status_code
        == 200
    )
    assert c.get(f"/api/contracts/{sale['id']}").json()["balance"] == 10_000_000
    assert (
        c.post(f"/api/payments/{paid['id']}/void", json={"reason": "Repetido"}).status_code == 409
    )
    assert c.get(f"/api/receipts/{paid['receipt_id']}/pdf").content.startswith(b"%PDF")
    paths = list(database.DATA.glob("recibos/*/*/*/*.pdf"))
    assert len(paths) == 1


def test_partial_payment_does_not_reset_arrears(portfolio):
    c, e, cl, lot, sale = portfolio
    before = c.get(f"/api/entities/{e['id']}/arrears").json()["items"][0]
    create(c, "/api/payments", payment(sale, "1"))
    after = c.get(f"/api/entities/{e['id']}/arrears").json()["items"][0]
    assert before["days_late"] == after["days_late"]
    assert after["overdue"] == before["overdue"] - 100
    assert after["color"] == "red"


def test_transfer_preserves_receipt(portfolio):
    c, e, cl, lot, sale = portfolio
    paid = create(c, "/api/payments", payment(sale))
    other = create(
        c,
        f"/api/entities/{e['id']}/clients",
        {
            "name": "Nuevo propietario sintético",
            "document_type": "DNI",
            "document": "SINTETICO-002",
            "phone": "000",
        },
    )
    today = c.get("/api/auth/me").json()["today"]
    response = c.post(
        f"/api/contracts/{sale['id']}/transfer",
        json={"new_client_id": other["id"], "effective_date": today, "reason": "Cesión de prueba"},
    )
    assert response.status_code == 200, response.text
    assert response.json()["balance"] == 9_000_000
    with database.transaction() as db:
        snapshot = json.loads(
            db.execute(
                "SELECT snapshot FROM receipts WHERE id=?", (paid["receipt_id"],)
            ).fetchone()[0]
        )
        assert snapshot["client"] == cl["name"]
        assert (
            db.execute("SELECT client_id FROM payments WHERE id=?", (paid["id"],)).fetchone()[0]
            == cl["id"]
        )


def test_operator_isolation_permissions_revocation(portfolio):
    c, e, cl, lot, sale = portfolio
    other = create(c, "/api/entities", {"name": "Lotificadora B", "currency": "HNL"})
    operator = create(
        c,
        "/api/users",
        {
            "username": "operador",
            "full_name": "Operador A",
            "password": "OperadorTemporal123",
            "entity_ids": [e["id"]],
        },
    )
    admin_payment = create(c, "/api/payments", payment(sale, "1"))
    with TestClient(app) as op:
        authenticate(op, "operador", "OperadorTemporal123")
        assert op.get("/api/entities").status_code == 403
        assert (
            op.post(
                "/api/auth/password",
                json={
                    "current_password": "OperadorTemporal123",
                    "new_password": "OperadorDefinitivo123",
                },
            ).status_code
            == 200
        )
        assert op.get(f"/api/entities/{other['id']}/clients").status_code == 403
        assert op.post("/api/entities", json={"name": "Escalada"}).status_code == 403
        assert op.get(f"/api/receipts/{admin_payment['receipt_id']}/pdf").status_code == 403
        paid = create(op, "/api/payments", payment(sale, "2"))
        assert (
            op.post(f"/api/payments/{paid['id']}/void", json={"reason": "No permitido"}).status_code
            == 403
        )
        items = op.get("/api/transactions").json()["items"]
        assert items and all(r["user_id"] == operator["id"] for r in items)
        assert len(op.get(f"/api/entities/{e['id']}/payments").json()["items"]) == 1
        approval = create(
            op,
            f"/api/entities/{e['id']}/approvals",
            {"action": "void_payment", "record_id": paid["id"], "reason": "Solicitud de prueba"},
        )
        assert (
            c.post(
                f"/api/approvals/{approval['id']}/resolve",
                json={"approve": True, "reason": "Aprobado en prueba"},
            ).status_code
            == 200
        )
        assert (
            c.put(
                f"/api/users/{operator['id']}",
                json={"full_name": "Operador A", "active": True, "entity_ids": []},
            ).status_code
            == 200
        )
        assert op.get(f"/api/receipts/{paid['receipt_id']}/pdf").status_code == 403
        assert op.get("/api/transactions").json()["items"] == []
        assert (
            c.post(
                f"/api/users/{operator['id']}/reset-password", json={"password": "NuevaTemporal123"}
            ).status_code
            == 200
        )
        assert op.get("/api/auth/me").status_code == 401


def test_csrf_logout_and_lockout(admin_client):
    c = admin_client
    csrf = c.headers.pop("x-csrf-token")
    assert c.post("/api/entities", json={"name": "No CSRF"}).status_code == 403
    c.headers["x-csrf-token"] = csrf
    assert (
        c.post(
            "/api/entities", json={"name": "Cross"}, headers={"origin": "https://evil.example"}
        ).status_code
        == 403
    )
    assert c.post("/api/auth/logout").status_code == 200
    assert c.get("/api/entities").status_code == 401
    for _ in range(5):
        assert (
            c.post(
                "/api/auth/login", json={"username": "admin", "password": "incorrecta"}
            ).status_code
            == 401
        )
    assert (
        c.post(
            "/api/auth/login", json={"username": "admin", "password": "ClaveSeguraPrueba123!"}
        ).status_code
        == 401
    )
    with database.transaction() as db:
        assert (
            db.execute("SELECT COUNT(*) FROM audit WHERE action='login_failed'").fetchone()[0] == 6
        )


def test_financial_terms_immutable_and_audit(portfolio):
    c, e, cl, lot, sale = portfolio
    assert (
        c.put(
            f"/api/lots/{lot['id']}",
            json={"code": lot["code"], "price": "200000", "currency": "HNL"},
        ).status_code
        == 200
    )
    assert c.get(f"/api/contracts/{sale['id']}").json()["price"] == 10_000_000
    with database.transaction() as db:
        try:
            db.execute("DELETE FROM audit")
        except sqlite3.IntegrityError:
            pass
        else:
            raise AssertionError("Auditoría borrada")


def test_concurrent_payment_limits(portfolio):
    c, e, cl, lot, sale = portfolio
    with database.transaction() as db:
        user = dict(db.execute("SELECT * FROM users WHERE role='superuser'").fetchone())

    def charge(index):
        try:
            with database.transaction() as db:
                return create_payment(db, user, Payment(**payment(sale, "60000")))
        except Exception as exc:
            return exc

    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(charge, range(2)))
    assert len([r for r in results if isinstance(r, dict)]) == 1
    with database.transaction() as db:
        assert db.execute("SELECT SUM(amount-paid) FROM obligations").fetchone()[0] == 4_000_000


def test_concurrent_sales(admin_client):
    c = admin_client
    e = create(c, "/api/entities", {"name": "Carrera"})
    cl = create(
        c,
        f"/api/entities/{e['id']}/clients",
        {"name": "Ejemplo", "document_type": "DNI", "document": "TEST", "phone": "000"},
    )
    lot = create(c, f"/api/entities/{e['id']}/lots", {"code": "UNO", "price": "100"})
    body = Sale(
        lot_id=lot["id"],
        client_id=cl["id"],
        price=Decimal(100),
        currency="HNL",
        down_payment=Decimal(0),
        installments=3,
        down_due=date(2026, 1, 31),
        first_due=date(2026, 1, 31),
    )
    with database.transaction() as db:
        user = dict(db.execute("SELECT * FROM users").fetchone())

    def sell(index):
        try:
            with database.transaction() as db:
                return create_sale(db, user, e["id"], body)
        except Exception as exc:
            return exc

    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(sell, range(2)))
    assert len([r for r in results if isinstance(r, dict)]) == 1
    with database.transaction() as db:
        assert [r[0] for r in db.execute("SELECT amount FROM obligations ORDER BY number")] == [
            0,
            3333,
            3333,
            3334,
        ]


def test_excel_backup_restore(portfolio):
    c, e, cl, lot, sale = portfolio
    create(c, "/api/payments", payment(sale))
    response = c.get("/api/export.xlsx")
    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.content))
    assert {"Pagos", "Clientes", "Lotes", "Contratos", "Cuotas"}.issubset(book.sheetnames)
    assert "password_hash" not in [cell.value for sheet in book for row in sheet for cell in row]
    result = create(c, "/api/backups", {})
    payload = c.get(f"/api/backups/{result['name']}").content
    create(
        c,
        f"/api/entities/{e['id']}/clients",
        {"name": "Posterior", "document_type": "DNI", "document": "POSTERIOR", "phone": "000"},
    )
    restored = c.post(
        "/api/backups/restore",
        files={"file": ("backup.zip", payload, "application/zip")},
        headers={"x-confirm-restore": "RESTAURAR"},
    )
    assert restored.status_code == 200, restored.text
    assert c.get("/api/auth/me").status_code == 401
    authenticate(c)
    assert c.get(f"/api/entities/{e['id']}/clients").json()["total"] == 1
    assert c.get(f"/api/contracts/{sale['id']}").json()["balance"] == 9_000_000
