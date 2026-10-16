import io
import json
from datetime import date
from uuid import uuid4

import pytest
from conftest import create
from openpyxl import load_workbook

from backend import db as database
from backend.finance import schedule


@pytest.fixture
def monthly(portfolio):
    c, e, client, *_ = portfolio
    lot = create(c, f"/api/entities/{e['id']}/lots", {"code": "CUOTA-2000", "price": "250000"})
    sale = create(
        c,
        f"/api/entities/{e['id']}/contracts",
        {
            "lot_id": lot["id"],
            "client_id": client["id"],
            "price": "250000",
            "currency": "HNL",
            "down_payment": "10000",
            "installments": 120,
            "down_due": "2026-09-25",
            "first_due": "2026-10-25",
        },
    )
    return c, e, sale


def body(sale, amount, **extra):
    return {
        "contract_id": sale["id"],
        "operation_id": str(uuid4()),
        "amount": amount,
        "currency": "HNL",
        **extra,
    }


def account(c, sale):
    return c.get(f"/api/contracts/{sale['id']}").json()


def amendment(a, **extra):
    return {
        "price": "250000",
        "down_payment": "10000",
        "installments": 120,
        "down_due": "2026-09-20",
        "first_due": "2026-10-20",
        "reason": "Corrección de fechas y condiciones",
        "expected_schedule": a["schedule_token"],
        **extra,
    }


def test_selected_month_overflow_and_partial_payment(monthly):
    c, e, sale = monthly
    create(c, "/api/payments", body(sale, "10000"))
    a = account(c, sale)
    data = body(sale, "2500", start_obligation_id=a["obligations"][1]["id"])
    preview = c.post("/api/payments/preview", json=data).json()
    assert [o["amount"] for o in preview["allocations"]] == [200000, 50000]
    assert [o["remaining_after"] for o in preview["allocations"]] == [0, 150000]
    data["allocation_token"] = preview["allocation_token"]
    p = create(c, "/api/payments", data)
    assert (
        p["months_paid"]
        == "Cuota 1 - octubre 2026: HNL 2,000.00; Cuota 2 - noviembre 2026: HNL 500.00 (abono parcial)"
    )
    assert create(c, "/api/payments", data)["id"] == p["id"]
    assert account(c, sale)["balance"] == 23750000
    next_payment = create(c, "/api/payments", body(sale, "500"))
    assert next_payment["allocations"][0]["number"] == 2
    assert next_payment["allocations"][0]["remaining_after"] == 100000
    with database.transaction() as db:
        snap = json.loads(
            db.execute("SELECT snapshot FROM receipts WHERE payment_id=?", (p["id"],)).fetchone()[0]
        )
        assert snap["allocations"][1]["remaining_after"] == 150000
    for url, sheet in (
        (f"/api/reports/export.xlsx?kind=payments&entity={e['id']}", "Cobros"),
        (f"/api/my-payments.xlsx?entity={e['id']}", "Mis pagos"),
    ):
        response = c.get(url)
        assert response.status_code == 200, response.text
        book = load_workbook(io.BytesIO(response.content))
        assert "noviembre 2026" in str(list(book[sheet].values))
        assert "Aplicaciones por mes" in book.sheetnames
    assert c.get(f"/api/receipts/{p['receipt_id']}/pdf").content.startswith(b"%PDF")


def test_selected_future_month_leaves_earlier_debt_and_limits_overflow(monthly):
    c, _, sale = monthly
    create(c, "/api/payments", body(sale, "10000"))
    a = account(c, sale)
    preview = c.post(
        "/api/payments/preview",
        json=body(sale, "2500", start_obligation_id=a["obligations"][3]["id"]),
    )
    assert preview.json()["prior_pending"] == 400000
    assert [o["number"] for o in preview.json()["allocations"]] == [3, 4]
    before = a["balance"]
    response = c.post(
        "/api/payments", json=body(sale, "2500", start_obligation_id=a["obligations"][-1]["id"])
    )
    assert response.status_code == 409
    assert account(c, sale)["balance"] == before
    assert (
        c.post("/api/payments", json=body(sale, "100", start_obligation_id=999999)).status_code
        == 409
    )


def test_conversion_and_stale_preview(monthly):
    c, _, sale = monthly
    create(c, "/api/payments", body(sale, "10000"))
    assert c.put("/api/settings", json={"exchange_rate": "25"}).status_code == 200
    data = body(sale, "100", currency="USD")
    preview = c.post("/api/payments/preview", json=data).json()
    assert preview["applied"] == 250000
    assert [a["amount"] for a in preview["allocations"]] == [200000, 50000]
    create(c, "/api/payments", body(sale, "100"))
    data["allocation_token"] = preview["allocation_token"]
    assert c.post("/api/payments", json=data).status_code == 409


def test_amend_in_place_retains_ids_paid_values_and_receipt_months(monthly):
    c, _, sale = monthly
    p = create(c, "/api/payments", body(sale, "14000"))
    before = account(c, sale)
    for due in ("2027-01-20", "2027-01-25"):
        data = amendment(account(c, sale), first_due=due)
        response = c.put(f"/api/contracts/{sale['id']}", json=data)
        assert response.status_code == 200, response.text
        after = account(c, sale)
        assert [o["id"] for o in after["obligations"]] == [o["id"] for o in before["obligations"]]
        assert len(after["obligations"]) == 121
        assert all(o["amount"] == 200000 for o in after["obligations"][1:])
        assert after["obligations"][1]["paid"] == after["obligations"][2]["paid"] == 200000
        assert after["balance"] == 23600000
        assert "octubre 2026" in after["payments"][0]["months_paid"]
    assert (
        c.post(f"/api/payments/{p['id']}/void", json={"reason": "Reversión comprobada"}).status_code
        == 200
    )
    assert account(c, sale)["balance"] == 25000000


def test_approval_amends_same_schedule_and_rejects_stale_request(monthly):
    c, e, sale = monthly
    create(c, "/api/payments", body(sale, "14000"))
    before = account(c, sale)
    data = amendment(before)
    reason = data.pop("reason")
    req = create(
        c,
        f"/api/entities/{e['id']}/approvals",
        {"action": "amend_contract", "record_id": sale["id"], "payload": data, "reason": reason},
    )
    assert account(c, sale)["version"] == before["version"]
    response = c.post(
        f"/api/approvals/{req['id']}/resolve", json={"approve": True, "reason": "Revisión aprobada"}
    )
    assert response.status_code == 200, response.text
    after = account(c, sale)
    assert [o["id"] for o in after["obligations"]] == [o["id"] for o in before["obligations"]]
    assert after["first_due"] == "2026-10-20"
    data = amendment(after)
    reason = data.pop("reason")
    req = create(
        c,
        f"/api/entities/{e['id']}/approvals",
        {"action": "amend_contract", "record_id": sale["id"], "payload": data, "reason": reason},
    )
    create(c, "/api/payments", body(sale, "100"))
    assert (
        c.post(
            f"/api/approvals/{req['id']}/resolve",
            json={"approve": True, "reason": "Se debe actualizar"},
        ).status_code
        == 409
    )
    assert account(c, sale)["version"] == after["version"]


@pytest.mark.parametrize("received", ["14000", "10500"])
def test_legacy_zeroed_schedule_is_compacted_without_erasing_payments(monthly, received):
    c, e, sale = monthly
    p = create(c, "/api/payments", body(sale, received))
    before = account(c, sale)
    with database.transaction() as db:
        db.execute(
            "INSERT INTO contract_versions(contract_id,version,snapshot,reason,created_by,created_at) VALUES(?,?,?,?,?,?)",
            (sale["id"], 1, json.dumps(before), "Versión anterior", 1, database.now()),
        )
        db.execute("UPDATE contracts SET version=2 WHERE id=?", (sale["id"],))
        db.execute(
            "UPDATE obligations SET amount=paid WHERE contract_id=? AND number>0", (sale["id"],)
        )
        schedule(
            db,
            sale["id"],
            25000000 - int(float(received) * 100),
            0,
            120,
            date(2026, 9, 25),
            date(2026, 10, 25),
            include_down=False,
            start_number=121,
        )
    assert len(account(c, sale)["obligations"]) == 241
    response = c.put(f"/api/contracts/{sale['id']}", json=amendment(account(c, sale)))
    assert response.status_code == 200, response.text
    after = account(c, sale)
    assert len(after["obligations"]) == 121
    assert all(o["amount"] == 200000 for o in after["obligations"][1:])
    assert after["balance"] == 25000000 - int(float(received) * 100)
    with database.transaction() as db:
        assert (
            db.execute(
                "SELECT COUNT(*) FROM obligations WHERE contract_id=? AND retired=1", (sale["id"],)
            ).fetchone()[0]
            == 120
        )
    report = load_workbook(io.BytesIO(c.get(f"/api/reports/export.xlsx?entity={e['id']}").content))
    assert sum(r[0] == sale["id"] for r in list(report["Plan de cuotas"].values)[1:]) == 121
    assert (
        c.post(f"/api/payments/{p['id']}/void", json={"reason": "Reversión conservada"}).status_code
        == 200
    )
    assert account(c, sale)["balance"] == 25000000


def test_partial_month_and_reducing_count_preserve_paid_applications(monthly):
    c, _, sale = monthly
    create(c, "/api/payments", body(sale, "10500"))
    before = account(c, sale)
    assert c.put(f"/api/contracts/{sale['id']}", json=amendment(before)).status_code == 200
    after = account(c, sale)
    assert after["obligations"][1]["paid"] == 50000
    assert all(o["amount"] == 200000 for o in after["obligations"][1:])
    create(c, "/api/payments", body(sale, "3500"))
    before = account(c, sale)
    response = c.put(f"/api/contracts/{sale['id']}", json=amendment(before, installments=1))
    assert response.status_code == 409
    assert account(c, sale)["version"] == before["version"]
