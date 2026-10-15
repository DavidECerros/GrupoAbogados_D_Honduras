import io
import uuid
from datetime import timedelta

from conftest import create
from openpyxl import load_workbook
from test_land_features import operator
from test_record_tools import call

from backend import db as database


def payment(c, sale, value, day):
    return create(
        c,
        "/api/payments",
        {
            "contract_id": sale["id"],
            "amount": value,
            "currency": "HNL",
            "payment_date": day,
            "method": "cash",
            "operation_id": str(uuid.uuid4()),
        },
    )


def test_purge_releases_code_without_reusing_identity(portfolio):
    c, e, *_ = portfolio
    lot = create(c, f"/api/entities/{e['id']}/lots", {"code": "LM-001", "price": "250000"})
    assert call(c, "lots", lot, "purge").status_code == 409
    assert call(c, "lots", lot).status_code == 200
    conflict = c.post(f"/api/entities/{e['id']}/lots", json={"code": "LM-001"})
    assert conflict.status_code == 409 and "eliminado" in conflict.text
    assert call(c, "lots", lot, "purge").status_code == 200
    new = create(c, f"/api/entities/{e['id']}/lots", {"code": "LM-001", "price": "250000"})
    assert new["id"] > lot["id"]
    assert call(c, "lots", lot, "restore").status_code == 404
    with database.transaction() as db:
        assert db.execute("SELECT COUNT(*) FROM audit WHERE action='purge_lots'").fetchone()[0] == 1


def test_purge_protects_financial_history(portfolio):
    c, _, customer, lot, sale = portfolio
    for kind, record in (("lots", lot), ("clients", customer)):
        assert call(c, kind, record).status_code == 200
        response = call(c, kind, record, "purge")
        assert response.status_code == 409 and "contracts" in response.text
    assert c.get(f"/api/contracts/{sale['id']}").status_code == 200


def test_fixed_price_is_master_only_and_new_sale_uses_saved_value(portfolio):
    c, e, customer, *_ = portfolio
    lot = create(c, f"/api/entities/{e['id']}/lots", {"code": "FIJO", "price": "250000"})
    for price in ("260000", "270000"):
        assert (
            c.put(f"/api/lots/{lot['id']}", json={"code": "FIJO", "price": price}).status_code
            == 200
        )
    operator(c, e, profile="administrator")
    for body in (
        {"code": "FIJO", "price": "1"},
        {"code": "FIJO", "price": "270000", "currency": "USD"},
        {"code": "FIJO", "price": "270000", "pricing_unit": "m2"},
    ):
        assert c.put(f"/api/lots/{lot['id']}", json=body).status_code == 403
    assert (
        c.put(
            f"/api/lots/{lot['id']}",
            json={"code": "FIJO", "price": "270000", "description": "Corrección permitida"},
        ).status_code
        == 200
    )
    body = {
        "lot_id": lot["id"],
        "client_id": customer["id"],
        "currency": "HNL",
        "price": "1",
        "down_payment": "0",
        "installments": 12,
        "down_due": "2026-01-01",
        "first_due": "2026-02-01",
    }
    assert c.post(f"/api/entities/{e['id']}/contracts", json=body).status_code == 403
    body["price"] = "270000"
    sale = create(c, f"/api/entities/{e['id']}/contracts", body)
    assert sale["price"] == 27000000
    assert call(c, "lots", lot, "purge").status_code == 403


def test_statement_exports_match_ledger_and_are_typed(portfolio):
    c, _, _, _, sale = portfolio
    response = c.get(f"/api/contracts/{sale['id']}/statement.xlsx")
    assert response.status_code == 200, response.text
    book = load_workbook(io.BytesIO(response.content))
    sheet = book.worksheets[0]
    header = next(r[0].row for r in sheet if r[0].value == "Obligación")
    rows = list(sheet.iter_rows(min_row=header + 1, max_row=sheet.max_row - 1, values_only=True))
    assert len(rows) == 13
    assert sum(r[2] for r in rows) == 100000
    assert sheet.cell(sheet.max_row, 5).value == 100000
    assert rows[0][1].year == 2026
    pdf = c.get(f"/api/contracts/{sale['id']}/statement.pdf")
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    assert "no-store" in pdf.headers["cache-control"]
    other = create(c, "/api/entities", {"name": "Otra"})
    operator(c, other, profile="administrator")
    for fmt in ("xlsx", "pdf"):
        assert c.get(f"/api/contracts/{sale['id']}/statement.{fmt}").status_code == 403


def test_daily_net_and_financial_export_reconcile(portfolio):
    c, e, _, _, sale = portfolio
    with database.transaction() as db:
        day = database.today(db)
    a = payment(c, sale, "100", day.isoformat())
    payment(c, sale, "200", day.isoformat())
    payment(c, sale, "300", (day - timedelta(days=1)).isoformat())
    assert (
        c.post(f"/api/payments/{a['id']}/void", json={"reason": "Prueba anulación"}).status_code
        == 200
    )
    data = c.get(f"/api/dashboard?entity={e['id']}").json()
    assert data["collection_today"] == [{"currency": "HNL", "amount": 20000}]
    response = c.get(
        f"/api/reports/export.xlsx?kind=finances&entity={e['id']}&since={day}&until={day}"
    )
    assert response.status_code == 200, response.text
    book = load_workbook(io.BytesIO(response.content))
    values = list(book["Finanzas"].values)
    row = dict(zip(values[0], values[1]))
    assert row["Cobros netos en período"] == 200
    assert row["Cobros anulados en período"] == 100
    assert row["Cartera pendiente actual"] == 99500
    assert c.get("/api/reports/export.xlsx?since=2026-12-31&until=2026-01-01").status_code == 422


def test_closing_status_follows_balance_and_deed_workflow(portfolio):
    c, e, _, _, sale = portfolio
    url = f"/api/contracts/{sale['id']}/closing-status"
    assert c.get(f"/api/contracts/{sale['id']}").json()["commercial_status"] == "pending_payment"
    assert c.put(url, json={"status": "deeded", "reason": "No pagado"}).status_code == 409
    with database.transaction() as db:
        day = database.today(db).isoformat()
    p = payment(c, sale, "100000", day)
    assert c.get(f"/api/contracts/{sale['id']}").json()["commercial_status"] == "paid"
    for status in ("pending_deed", "deeded"):
        assert (
            c.put(url, json={"status": status, "reason": "Trámite verificado"}).status_code == 200
        )
        assert (
            c.get(f"/api/entities/{e['id']}/contracts").json()["items"][0]["commercial_status"]
            == status
        )
    assert (
        c.post(f"/api/payments/{p['id']}/void", json={"reason": "Corrección cobro"}).status_code
        == 200
    )
    assert c.get(f"/api/contracts/{sale['id']}").json()["commercial_status"] == "pending_payment"
    operator(c, e, profile="administrator")
    assert c.put(url, json={"status": "deeded", "reason": "No autorizado"}).status_code == 403


def test_report_types_entity_state_and_access(portfolio):
    c, e, *_ = portfolio
    other = create(c, "/api/entities", {"name": "Entidad B"})
    create(c, f"/api/entities/{other['id']}/lots", {"code": "SOLO-B", "price": "1"})
    for kind in ("all", "contracts", "payments", "lots", "finances", "overdue"):
        r = c.get(f"/api/reports/export.xlsx?kind={kind}&entity={e['id']}")
        assert r.status_code == 200, r.text
        book = load_workbook(io.BytesIO(r.content))
        assert "Entidad B" not in str([list(s.values) for s in book])
        assert "password_hash" not in str([list(s.values) for s in book])
    r = c.get(f"/api/reports/export.xlsx?kind=lots&state=available&entity={e['id']}")
    assert list(load_workbook(io.BytesIO(r.content))["Lotes"].values) == [("Sin registros",)]
    operator(c, e, profile="administrator")
    assert c.get(f"/api/reports/export.xlsx?entity={e['id']}").status_code == 403
