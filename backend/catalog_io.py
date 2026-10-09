"""Spreadsheet interchange for the two editable catalogs."""

import csv
import hashlib
import io
import json
import zipfile
from decimal import Decimal

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from pydantic import ValidationError

from . import models
from .db import audit, now
from .finance import cents, require

FIELDS = {
    "clients": [
        ("id", "ID"),
        ("name", "Nombre"),
        ("document_type", "Tipo de documento"),
        ("document", "Documento"),
        ("phone", "Teléfono"),
        ("email", "Correo"),
        ("notes", "Observaciones"),
        ("status", "Estado"),
    ],
    "lots": [
        ("id", "ID"),
        ("code", "Código"),
        ("description", "Descripción"),
        ("price", "Precio"),
        ("currency", "Moneda"),
        ("status", "Estado"),
    ],
}
STATUS = {
    "Activo": "active",
    "Inhabilitado": "inactive",
    "Disponible": "available",
    "Reservado": "reserved",
    "Vendido": "sold",
}


def kind_valid(kind):
    require(kind in FIELDS, "Catálogo inválido", 404)


def fingerprint(db, entity, kind):
    payload = [
        dict(r)
        for r in db.execute(f"SELECT * FROM {kind} WHERE entity_id=? ORDER BY id", (entity,))
    ]
    return hashlib.sha256(json.dumps(payload, sort_keys=True).encode()).hexdigest()


def export_catalog(db, entity, kind, q="", state=""):
    kind_valid(kind)
    search = (
        "(name LIKE ? OR document LIKE ?)"
        if kind == "clients"
        else "(code LIKE ? OR description LIKE ?)"
    )
    rows = db.execute(
        f"SELECT * FROM {kind} WHERE entity_id=? AND {search} AND (?='' OR status=?) ORDER BY id",
        (entity, f"%{q}%", f"%{q}%", state, state),
    ).fetchall()
    book = Workbook()
    sheet = book.active
    sheet.title = "Clientes" if kind == "clients" else "Lotes"
    sheet.append([title for field, title in FIELDS[kind]])
    for record in rows:
        values = []
        for field, title in FIELDS[kind]:
            value = record[field]
            if field == "price":
                value = Decimal(value) / 100
            if field == "status":
                value = next((label for label, code in STATUS.items() if code == value), value)
            values.append(value)
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"  # Keep literal strings, including leading '=', as text.
                cell.number_format = "@"
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17313B")
        cell.font = Font(color="FFFFFF", bold=True)
    for index, (field, title) in enumerate(FIELDS[kind], 1):
        from openpyxl.utils import get_column_letter

        sheet.column_dimensions[get_column_letter(index)].width = (
            36 if field in ("name", "description", "notes") else 24
        )
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def read_file(content, filename, kind):
    kind_valid(kind)
    require(len(content) <= 12_000_000, "Archivo máximo 12 MB")
    if filename.lower().endswith(".xlsx"):
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                require(
                    sum(item.file_size for item in archive.infolist()) <= 100_000_000,
                    "Excel descomprimido demasiado grande",
                )
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
            sheet = book["Clientes" if kind == "clients" else "Lotes"]
            records = []
            for row in sheet.iter_rows():
                require(len(records) <= 10000, "Máximo 10,000 filas")
                require(
                    not any(cell.data_type == "f" for cell in row),
                    "No se admiten fórmulas en la hoja de datos",
                )
                records.append([cell.value for cell in row])
            book.close()
        except (zipfile.BadZipFile, KeyError, ValueError, OSError):
            require(False, "Excel inválido. Use la hoja Clientes o Lotes de la plantilla")
    else:
        require(filename.lower().endswith(".csv"), "Use .xlsx o .csv")
        try:
            text = content.decode("utf-8-sig")
            delimiter = (
                ";" if text.splitlines()[0].count(";") > text.splitlines()[0].count(",") else ","
            )
            records = list(csv.reader(io.StringIO(text), delimiter=delimiter))
        except (UnicodeError, IndexError, csv.Error):
            require(False, "CSV inválido; guárdelo como UTF-8")
    require(records and len(records) <= 10001, "Archivo vacío o más de 10,000 filas")
    headers = [str(value or "").strip() for value in records[0]]
    expected = dict((title, field) for field, title in FIELDS[kind])
    require(len(headers) == len(set(headers)), "Encabezados duplicados")
    require(set(headers) == set(expected), "Los encabezados deben coincidir con la plantilla")
    return [
        (
            index,
            {
                expected[title]: values[col] if col < len(values) else None
                for col, title in enumerate(headers)
            },
        )
        for index, values in enumerate(records[1:], 2)
        if any(value not in (None, "") for value in values)
    ]


def plan_import(db, entity, kind, records, mode):
    require(mode in ("create", "upsert"), "Modo de importación inválido")
    errors, changes, seen = [], [], set()
    for line, values in records:
        try:
            values = dict(values)
            identifier = values.pop("id")
            require(
                identifier in (None, "") or (str(identifier).isdigit() and int(identifier) > 0),
                "ID debe ser un entero positivo o estar vacío",
            )
            identifier = int(identifier) if identifier not in (None, "") else None
            raw_status = str(values.pop("status") or "").strip()
            status = STATUS.get(raw_status, raw_status)
            for key, value in list(values.items()):
                if key != "price":
                    require(
                        not isinstance(value, (int, float)),
                        f"{key}: use texto para conservar ceros iniciales",
                    )
                    values[key] = str(value or "").strip()
            data = (models.Client if kind == "clients" else models.Lot)(**values).model_dump(
                mode="json"
            )
            natural = (
                (data["document_type"], data["document"]) if kind == "clients" else (data["code"],)
            )
            key_sql = "document_type=? AND document=?" if kind == "clients" else "code=?"
            previous = db.execute(
                f"SELECT * FROM {kind} WHERE entity_id=? AND {key_sql}", (entity, *natural)
            ).fetchone()
            if identifier:
                identified = db.execute(
                    f"SELECT * FROM {kind} WHERE id=? AND entity_id=?", (identifier, entity)
                ).fetchone()
                require(identified is not None, "ID no pertenece a esta lotificadora")
                require(
                    previous is None or previous["id"] == identifier, "Documento o código duplicado"
                )
                previous = identified
            require(
                mode == "upsert" or previous is None,
                "Registro existente; seleccione Crear y actualizar",
            )
            record_key = ("id", previous["id"]) if previous else natural
            require(
                record_key not in seen and natural not in seen, "Registro repetido en el archivo"
            )
            seen.add(record_key)
            seen.add(natural)
            if kind == "clients":
                status = status or (previous["status"] if previous else "active")
                require(status in ("active", "inactive"), "Estado debe ser Activo o Inhabilitado")
                data["status"] = status
            else:
                require(
                    not status or status == (previous["status"] if previous else "available"),
                    "El estado del lote depende de reservas y contratos; no se importa",
                )
            changes.append(
                {
                    "line": line,
                    "operation": "update" if previous else "create",
                    "id": previous["id"] if previous else None,
                    "data": data,
                }
            )
        except (ValidationError, ValueError, TypeError) as exc:
            errors.append(
                {"line": line, "message": str(exc).split("For further information")[0][:600]}
            )
        except Exception as exc:
            from fastapi import HTTPException

            if not isinstance(exc, HTTPException):
                raise
            errors.append({"line": line, "message": str(exc.detail)})
    require(records, "No hay filas para importar")
    return changes, errors


def apply_import(db, user, entity, kind, changes, reason):
    for item in changes:
        data, identifier = dict(item["data"]), item["id"]
        if kind == "lots":
            data["price"] = cents(data["price"])
        fields = list(data)
        if identifier:
            before = dict(db.execute(f"SELECT * FROM {kind} WHERE id=?", (identifier,)).fetchone())
            db.execute(
                f"UPDATE {kind} SET " + ",".join(f"{key}=?" for key in fields) + " WHERE id=?",
                (*data.values(), identifier),
            )
        else:
            before = None
            data["entity_id"] = entity
            if kind == "clients":
                data.update(created_by=user["id"], created_at=now())
            identifier = db.execute(
                f"INSERT INTO {kind}("
                + ",".join(data)
                + ") VALUES("
                + ",".join("?" for _ in data)
                + ")",
                tuple(data.values()),
            ).lastrowid
        after = dict(db.execute(f"SELECT * FROM {kind} WHERE id=?", (identifier,)).fetchone())
        audit(
            db,
            user,
            "import_" + kind,
            entity,
            kind,
            identifier,
            reason=reason,
            before=before,
            after=after,
        )
