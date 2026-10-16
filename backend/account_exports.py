"""Printable account and payment schedule, generated from the financial ledger."""

import io
from datetime import date
from decimal import Decimal
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

HEADERS = ["Obligación", "Vencimiento", "Importe", "Pagado", "Pendiente", "Situación"]


def situation(o, cutoff):
    if o["amount"] == 0:
        return "Sin importe"
    return (
        "Pagada"
        if o["paid"] == o["amount"]
        else "Vencida"
        if o["due_date"] < cutoff
        else "Pendiente"
    )


def export_account(account, entity, cutoff, format):
    currency = account["currency"]
    amounts = [sum(o[k] for o in account["obligations"]) for k in ("amount", "paid")]
    fields = [
        ("Lotificadora", entity["name"]),
        (
            "Contrato / versión / estado",
            f"#{account['id']} / {account['version']} / "
            + ("Activo" if account["status"] == "active" else "Anulado"),
        ),
        ("Cliente actual", account["client"]["name"]),
        ("Documento", account["client"]["document"]),
        ("Lote", account["lot"]["code"]),
        ("Fecha de corte", cutoff),
        ("Moneda", currency),
        (
            "Estado de cancelación / escritura",
            {
                "pending_payment": "Pendiente de cancelar",
                "paid": "Cancelado (pagado)",
                "pending_deed": "Pendiente de escriturar",
                "deeded": "Escriturado",
                "annulled": "Anulado",
            }[account["commercial_status"]],
        ),
        ("Precio del contrato", account["price"]),
        ("Prima acordada", account["down_payment"]),
        ("Monto original a financiar", account["price"] - account["down_payment"]),
        ("Importe del plan vigente", amounts[0]),
        ("Pagado aplicado al plan", amounts[1]),
        ("Saldo pendiente", account["balance"]),
    ]
    rows = [
        [
            "Prima" if o["number"] == 0 else f"Cuota {o['number']}",
            o["due_date"],
            o["amount"],
            o["paid"],
            o["amount"] - o["paid"],
            situation(o, cutoff),
        ]
        for o in account["obligations"]
    ]
    note = "Plan vigente a la fecha de corte. La prima acordada no acredita un pago. Los cobros anulados no reducen el saldo. No es recibo ni factura."
    if format == "xlsx":
        book = Workbook()
        sheet = book.active
        sheet.title = "Estado y plan de pagos"
        sheet.append(["ESTADO DE CUENTA Y PLAN DE PAGOS"])
        sheet.merge_cells("A1:F1")
        for label, value in fields:
            sheet.append([label, Decimal(value) / 100 if isinstance(value, int) else value])
            if not isinstance(value, int):
                sheet.merge_cells(
                    start_row=sheet.max_row, start_column=2, end_row=sheet.max_row, end_column=6
                )
            sheet.cell(sheet.max_row, 1).alignment = Alignment(wrap_text=True, vertical="center")
            sheet.cell(sheet.max_row, 2).alignment = Alignment(wrap_text=True, vertical="center")
            sheet.row_dimensions[sheet.max_row].height = 30
            if isinstance(value, int):
                sheet.cell(sheet.max_row, 2).number_format = f'"{currency} "#,##0.00'
        sheet.append([note])
        sheet.merge_cells(
            start_row=sheet.max_row, start_column=1, end_row=sheet.max_row, end_column=6
        )
        sheet.row_dimensions[sheet.max_row].height = 32
        sheet.cell(sheet.max_row, 1).alignment = Alignment(wrap_text=True, vertical="center")
        sheet.append([])
        header = sheet.max_row + 1
        sheet.append(HEADERS)
        for row in rows:
            sheet.append(
                [row[0], date.fromisoformat(row[1]), *[Decimal(v) / 100 for v in row[2:5]], row[5]]
            )
            sheet.cell(sheet.max_row, 2).number_format = "dd/mm/yyyy"
            for col in (3, 4, 5):
                sheet.cell(sheet.max_row, col).number_format = f'"{currency} "#,##0.00'
        sheet.append(
            [
                "TOTAL",
                None,
                Decimal(amounts[0]) / 100,
                Decimal(amounts[1]) / 100,
                Decimal(account["balance"]) / 100,
            ]
        )
        for cell in sheet[sheet.max_row][2:5]:
            cell.number_format = f'"{currency} "#,##0.00'
        payments = book.create_sheet("Cobros")
        payments.append(
            [
                "Fecha",
                "Folio",
                "Recibido",
                "Moneda recibida",
                "Aplicado al contrato",
                "Estado",
                "Meses / cuotas pagados",
            ]
        )
        for p in account["payments"]:
            payments.append(
                [
                    date.fromisoformat(p["payment_date"]),
                    p["folio"],
                    Decimal(p["received"]) / 100,
                    p["received_currency"],
                    Decimal(p["applied"]) / 100,
                    "Confirmado" if p["status"] == "valid" else "Anulado",
                    p.get("months_paid", ""),
                ]
            )
            payments.cell(payments.max_row, 1).number_format = "dd/mm/yyyy"
            payments.cell(
                payments.max_row, 3
            ).number_format = f'"{p["received_currency"]} "#,##0.00'
            payments.cell(payments.max_row, 5).number_format = f'"{currency} "#,##0.00'
        for ws, head, widths in (
            (sheet, header, [33, 28, 22, 22, 22, 20]),
            (payments, 1, [18, 30, 22, 23, 27, 20, 65]),
        ):
            ws.sheet_view.showGridLines = False
            ws.freeze_panes = f"C{head + 1}"
            ws.auto_filter.ref = (
                f"A{head}:{'F' if ws == sheet else 'G'}{ws.max_row - (1 if ws == sheet else 0)}"
            )
            for i, width in enumerate(widths, 1):
                from openpyxl.utils import get_column_letter

                ws.column_dimensions[get_column_letter(i)].width = width
            for r in ws:
                for cell in r:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"  # User text is never an Excel formula.
                    if cell.row in (1, head) or (ws == sheet and cell.row == ws.max_row):
                        cell.fill = PatternFill("solid", fgColor="17313B")
                        cell.font = Font(bold=True, color="FFFFFF")
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_setup.orientation = "landscape"
            ws.page_setup.paperSize = ws.PAPERSIZE_A4
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.print_title_rows = f"1:{head}"
            ws.print_options.horizontalCentered = True
        stream = io.BytesIO()
        book.save(stream)
        return stream.getvalue()

    styles = getSampleStyleSheet()
    normal = styles["Normal"]
    normal.fontSize = 9
    normal.leading = 12

    def para(value):
        return Paragraph(escape(str(value)), normal)

    def money(value):
        return f"{currency} {value / 100:,.2f}"

    contents = [Paragraph("Estado de cuenta y plan de pagos", styles["Title"])]
    summary = Table(
        [[para(k), para(money(v) if isinstance(v, int) else v)] for k, v in fields],
        colWidths=[190, 350],
        hAlign="LEFT",
    )
    summary.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.whitesmoke, colors.white]),
            ]
        )
    )
    contents += [summary, Spacer(1, 12), para(note), Spacer(1, 12)]
    table_rows = [HEADERS] + [[r[0], r[1], *[money(v) for v in r[2:5]], r[5]] for r in rows]
    table_rows += [
        ["TOTAL", "", money(amounts[0]), money(amounts[1]), money(account["balance"]), ""]
    ]
    table = Table(
        [[para(c) for c in r] for r in table_rows],
        colWidths=[64, 76, 105, 105, 105, 85],
        repeatRows=1,
    )
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E1ECE8")),
                ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E1ECE8")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.whitesmoke, colors.white]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    contents.append(table)
    if account["payments"]:
        contents += [Spacer(1, 15), Paragraph("Cobros registrados", styles["Heading2"])]
        payment_rows = [["Fecha / folio", "Recibido", "Aplicado", "Estado"]]
        for p in account["payments"]:
            payment_rows.append(
                [
                    p["payment_date"] + " / " + p["folio"] + " - " + p.get("months_paid", ""),
                    f"{p['received_currency']} {p['received'] / 100:,.2f}",
                    money(p["applied"]),
                    "Confirmado" if p["status"] == "valid" else "Anulado",
                ]
            )
        payment_table = Table(
            [[para(c) for c in r] for r in payment_rows],
            colWidths=[215, 115, 115, 95],
            repeatRows=1,
        )
        payment_table.setStyle(
            TableStyle(
                [
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E1ECE8")),
                ]
            )
        )
        contents.append(payment_table)
    stream = io.BytesIO()

    def footer(canvas, doc):
        canvas.setFont("Helvetica", 8)
        canvas.drawString(36, 20, f"Contrato #{account['id']} - Corte {cutoff}")
        canvas.drawRightString(576, 20, f"Página {doc.page}")

    SimpleDocTemplate(
        stream, pagesize=letter, leftMargin=36, rightMargin=36, topMargin=30, bottomMargin=38
    ).build(contents, onFirstPage=footer, onLaterPages=footer)
    return stream.getvalue()
