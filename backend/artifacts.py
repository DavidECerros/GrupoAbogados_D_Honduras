import hashlib
import io
import json
import os
import shutil
import sqlite3
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from . import db as database
from .db import audit, now, setting, today
from .finance import require, row


def generate_receipt(db, identifier):
    receipt = row(db, "receipts", identifier)
    snap = json.loads(receipt["snapshot"])
    issued = datetime.fromisoformat(snap["issued_date"])
    relative = (
        Path("recibos")
        / str(receipt["entity_id"])
        / str(receipt["year"])
        / f"{issued.month:02}"
        / f"{receipt['folio']}.pdf"
    )
    path = database.DATA / relative
    path.parent.mkdir(parents=True, exist_ok=True)
    styles = getSampleStyleSheet()
    page = letter if setting(db, "paper") == "letter" else (letter[0], letter[1] / 2)
    contents = []
    if snap.get("logo") and (database.DATA / snap["logo"]).is_file():
        contents.append(
            Image(str(database.DATA / snap["logo"]), width=80, height=40, kind="proportional")
        )
    contents += [
        Paragraph(escape(snap["entity"]), styles["Title"]),
        Paragraph(
            escape(snap["folio"]) + (" · ANULADO" if receipt["status"] == "void" else ""),
            styles["Heading2"],
        ),
        Paragraph("Recibo de cobro · No es factura fiscal", styles["Normal"]),
        Spacer(1, 12),
    ]

    def money(value, currency):
        return f"{currency} {value / 100:,.2f}"

    fields = [
        ("Entidad ID", str(snap["entity_id"])),
        ("Cliente", snap["client"]),
        ("Lote", snap["lot"]),
        ("Fecha de pago / emisión", f"{snap['payment_date']} / {snap['issued_date']}"),
        ("Importe recibido", money(snap["received"], snap["received_currency"])),
        ("Aplicado al contrato", money(snap["applied"], snap["contract_currency"])),
        ("Tasa HNL por USD", snap["rate"] or "No aplica"),
        (
            "Método / comprobante",
            f"{ {'cash': 'Efectivo', 'deposit': 'Depósito', 'transfer': 'Transferencia'}.get(snap['method'], snap['method']) } / {snap['bank_reference'] or '-'}",
        ),
        ("Responsable", snap["operator"]),
        ("Saldo posterior al cobro", money(snap["balance_after"], snap["contract_currency"])),
    ]
    if receipt["status"] == "void":
        payment = row(db, "payments", receipt["payment_id"])
        fields += [
            ("Anulación", payment["void_reason"]),
            ("Fecha de anulación", payment["void_at"]),
        ]
    table = Table(
        [
            [
                Paragraph(escape(str(a)), styles["Normal"]),
                Paragraph(escape(str(b)), styles["Normal"]),
            ]
            for a, b in fields
        ],
        colWidths=[160, page[0] - 220],
        hAlign="LEFT",
    )
    table.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.whitesmoke, colors.white]),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    contents.append(table)
    temporary = path.with_suffix(".tmp")
    SimpleDocTemplate(
        str(temporary), pagesize=page, leftMargin=30, rightMargin=30, topMargin=20, bottomMargin=20
    ).build(contents)
    os.replace(temporary, path)
    db.execute(
        "UPDATE receipts SET path=?,status=? WHERE id=?",
        (relative.as_posix(), "void" if receipt["status"] == "void" else "issued", identifier),
    )
    return path


def export_excel(db, tables):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, records in tables.items():
        sheet = workbook.create_sheet(name[:31])
        if records:
            sheet.append(list(records[0]))
            for record in records:
                # Escape formula-like strings so notes/documents cannot execute formulas.
                sheet.append(
                    [
                        "'" + v if isinstance(v, str) and v.startswith(("=", "+", "-", "@")) else v
                        for v in record.values()
                    ]
                )
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="18384A")
            for column in sheet.columns:
                sheet.column_dimensions[column[0].column_letter].width = min(
                    45, max(14, max(len(str(c.value or "")) for c in column[:100]) + 2)
                )
        else:
            sheet.append(["Sin registros"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def backup(user=None, daily=False):
    with database.LOCK:
        db = database.connect()
        destination = Path(setting(db, "backup_destination") or str(database.DATA / "backups"))
        destination.mkdir(parents=True, exist_ok=True)
        day = today(db).isoformat()
        name = (
            f"daily-{day}.zip"
            if daily
            else f"manual-{datetime.now().strftime('%Y%m%d-%H%M%S-%f')}.zip"
        )
        target = destination / name
        if daily and target.exists():
            db.close()
            return target
        try:
            with tempfile.TemporaryDirectory() as temporary:
                snapshot = Path(temporary) / "gestor.sqlite3"
                copy = sqlite3.connect(snapshot)
                db.backup(copy)
                copy.close()
                files = {"gestor.sqlite3": snapshot}
                for directory in ("recibos", "logos"):
                    for path in (database.DATA / directory).rglob("*"):
                        if path.is_file() and path.suffix != ".tmp":
                            files[path.relative_to(database.DATA).as_posix()] = path
                manifest = {
                    "version": 1,
                    "created_at": now(),
                    "files": {name: file_digest(path) for name, path in files.items()},
                }
                pending = target.with_suffix(".tmp")
                with zipfile.ZipFile(pending, "w", zipfile.ZIP_DEFLATED) as archive:
                    archive.writestr("manifest.json", json.dumps(manifest))
                    for name, path in files.items():
                        archive.write(path, name)
                os.replace(pending, target)
            db.execute("BEGIN IMMEDIATE")
            audit(
                db,
                user,
                "daily_backup" if daily else "manual_backup",
                kind="backup",
                after={"name": target.name},
            )
            db.commit()
            if daily:
                for old in sorted(destination.glob("daily-*.zip"))[:-30]:
                    old.unlink()
            return target
        finally:
            db.close()


def file_digest(path):
    with path.open("rb") as file:
        return hashlib.file_digest(file, "sha256").hexdigest()


def restore(content, user):
    with database.LOCK, tempfile.TemporaryDirectory(dir=database.DATA) as temporary:
        staged = Path(temporary)
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                require(
                    sum(i.file_size for i in archive.infolist()) <= 2_000_000_000,
                    "Respaldo demasiado grande",
                )
                names = archive.namelist()
                require(len(names) == len(set(names)), "Entradas duplicadas en respaldo")
                manifest = json.loads(archive.read("manifest.json"))
                require(manifest["version"] == 1, "Versión incompatible")
                require(
                    set(names) == set(manifest["files"]) | {"manifest.json"},
                    "Contenido no declarado",
                )
                for name, digest in manifest["files"].items():
                    relative = Path(name)
                    require(
                        not relative.is_absolute()
                        and ".." not in relative.parts
                        and "\\" not in name
                        and ":" not in name,
                        "Ruta de respaldo inválida",
                    )
                    require(
                        name == "gestor.sqlite3" or name.startswith(("recibos/", "logos/")),
                        "Archivo no permitido",
                    )
                    payload = archive.read(name)
                    require(
                        hashlib.sha256(payload).hexdigest() == digest,
                        "Respaldo alterado o corrupto",
                    )
                    target = staged / relative
                    target.parent.mkdir(parents=True, exist_ok=True)
                    target.write_bytes(payload)
            require((staged / "gestor.sqlite3").is_file(), "Falta la base de datos")
            probe = sqlite3.connect(staged / "gestor.sqlite3")
            try:
                require(
                    probe.execute("PRAGMA integrity_check").fetchone()[0] == "ok",
                    "Base de datos corrupta",
                )
                require(
                    not probe.execute("PRAGMA foreign_key_check").fetchall(), "Relaciones inválidas"
                )
                require(
                    probe.execute("SELECT version FROM schema_version").fetchone()[0] == 1,
                    "Esquema incompatible",
                )
                superuser = probe.execute(
                    "SELECT id,role FROM users WHERE role='superuser' AND active=1"
                ).fetchall()
                require(len(superuser) == 1, "Respaldo sin superusuario único")
            finally:
                probe.close()
        except (zipfile.BadZipFile, KeyError, ValueError, sqlite3.DatabaseError):
            require(False, "Respaldo inválido")
        # Caller is the sole active session and no application transaction is open.
        previous = backup(user)
        current = database.connect()
        replacement = sqlite3.connect(staged / "gestor.sqlite3")
        rollback_db = sqlite3.connect(staged / "before.sqlite3")
        old_folders = []
        new_folders = []
        try:
            current.backup(rollback_db)
            for folder in ("recibos", "logos"):
                target = database.DATA / folder
                if target.exists():
                    os.replace(target, staged / ("before-" + folder))
                    old_folders.append(folder)
            replacement.backup(current)
            for folder in ("recibos", "logos"):
                if (staged / folder).exists():
                    os.replace(staged / folder, database.DATA / folder)
                    new_folders.append(folder)
        except (OSError, sqlite3.DatabaseError):
            rollback_db.backup(current)
            for folder in new_folders:
                target = database.DATA / folder
                if target.exists():
                    shutil.rmtree(target)
            for folder in old_folders:
                os.replace(staged / ("before-" + folder), database.DATA / folder)
            audit(
                current,
                user,
                "restore_failed",
                kind="backup",
                reason="No se pudo completar el reemplazo; se revirtió al estado previo",
                result="failed",
            )
            raise
        finally:
            current.close()
            replacement.close()
            rollback_db.close()
        with database.transaction() as db:
            db.execute("DELETE FROM sessions")
            restored_user = dict(
                db.execute("SELECT id,role FROM users WHERE role='superuser'").fetchone()
            )
            audit(
                db,
                restored_user,
                "restore",
                reason="Restauración local autorizada",
                after={"previous_backup": previous.name, "requested_by_username": user["username"]},
            )
        return {"restored": True, "previous_backup": previous.name, "login_required": True}
